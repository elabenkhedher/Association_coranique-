"""
Views pour l'application statistiques.
Calculs globaux et par entité, exports PDF et Excel.
"""
import io
from datetime import date

from rest_framework.views import APIView
from rest_framework.response import Response
from rest_framework.permissions import IsAuthenticated
from django.http import HttpResponse
from django.db.models import Count, Avg

from users.models import CustomUser
from centres.models import Centre
from planning.models import Creneau
from presences.models import Presence
from lectures.models import SeanceLecture
from users.permissions import EstAdmin


class StatistiquesGlobalesView(APIView):
    """
    GET /api/statistiques/global/
    Statistiques globales : centres, enseignants, apprenants, séances.
    Réservé aux administrateurs.
    """
    permission_classes = [EstAdmin]

    def get(self, request):
        return Response({
            'nombre_centres': Centre.objects.count(),
            'nombre_enseignants': CustomUser.objects.filter(role='enseignant').count(),
            'nombre_apprenants': CustomUser.objects.filter(role='apprenant').count(),
            'nombre_seances_lecture': SeanceLecture.objects.count(),
            'nombre_creneaux': Creneau.objects.count(),
            'creneaux_annules': Creneau.objects.filter(est_annule=True).count(),
        })


class StatistiquesCentreView(APIView):
    """
    GET /api/statistiques/centre/{id}/
    Stats détaillées d'un centre.
    """
    permission_classes = [IsAuthenticated]

    def get(self, request, centre_id):
        try:
            centre = Centre.objects.get(pk=centre_id)
        except Centre.DoesNotExist:
            return Response({'erreur': 'Centre introuvable.'}, status=404)

        # Restriction par rôle
        if request.user.role != 'admin' and request.user.centre_id != centre_id:
            return Response({'erreur': 'Accès refusé.'}, status=403)

        # Taux de présence
        total = Presence.objects.filter(creneau__centre=centre).count()
        presents = Presence.objects.filter(creneau__centre=centre, statut='present').count()
        taux_presence = round(presents / total * 100, 2) if total > 0 else 0

        # Progression moyenne des apprenants du centre
        seances = SeanceLecture.objects.filter(enseignant__centre=centre)
        progression_moy = seances.aggregate(Avg('evaluation'))['evaluation__avg'] or 0

        return Response({
            'centre': centre.nom,
            'nombre_enseignants': CustomUser.objects.filter(centre=centre, role='enseignant').count(),
            'nombre_apprenants': CustomUser.objects.filter(centre=centre, role='apprenant').count(),
            'nombre_creneaux': Creneau.objects.filter(centre=centre).count(),
            'nombre_seances': seances.count(),
            'taux_presence': taux_presence,
            'evaluation_moyenne': round(float(progression_moy), 2),
        })


class StatistiquesEnseignantView(APIView):
    """
    GET /api/statistiques/enseignant/{id}/
    Stats d'un enseignant : séances enregistrées, apprenants suivis.
    """
    permission_classes = [IsAuthenticated]

    def get(self, request, enseignant_id):
        try:
            enseignant = CustomUser.objects.get(pk=enseignant_id, role='enseignant')
        except CustomUser.DoesNotExist:
            return Response({'erreur': 'Enseignant introuvable.'}, status=404)

        seances = SeanceLecture.objects.filter(enseignant=enseignant)
        apprenants_suivis = seances.values('apprenant').distinct().count()

        return Response({
            'enseignant': str(enseignant),
            'nombre_seances': seances.count(),
            'apprenants_suivis': apprenants_suivis,
            'evaluation_moyenne': round(
                float(seances.aggregate(Avg('evaluation'))['evaluation__avg'] or 0), 2
            ),
        })


class StatistiquesApprenantView(APIView):
    """
    GET /api/statistiques/apprenant/{id}/
    Stats d'un apprenant : progression, taux de présence, séances.
    """
    permission_classes = [IsAuthenticated]

    def get(self, request, apprenant_id):
        # Restriction par rôle
        if request.user.role == 'apprenant' and request.user.id != apprenant_id:
            return Response({'erreur': 'Accès refusé.'}, status=403)

        try:
            apprenant = CustomUser.objects.get(pk=apprenant_id, role='apprenant')
        except CustomUser.DoesNotExist:
            return Response({'erreur': 'Apprenant introuvable.'}, status=404)

        # Présences
        presences = Presence.objects.filter(apprenant=apprenant)
        total = presences.count()
        presents = presences.filter(statut='present').count()
        taux_presence = round(presents / total * 100, 2) if total > 0 else 0

        # Séances de lecture
        seances = SeanceLecture.objects.filter(apprenant=apprenant)
        eval_moy = seances.aggregate(Avg('evaluation'))['evaluation__avg'] or 0

        return Response({
            'apprenant': str(apprenant),
            'nombre_seances': seances.count(),
            'taux_presence': taux_presence,
            'evaluation_moyenne': round(float(eval_moy), 2),
        })


class ExportPDFView(APIView):
    """
    GET /api/statistiques/export-pdf/
    Export PDF du rapport global avec reportlab.
    """
    permission_classes = [EstAdmin]

    def get(self, request):
        try:
            from reportlab.lib.pagesizes import A4
            from reportlab.lib.styles import getSampleStyleSheet
            from reportlab.lib.units import cm
            from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle, Spacer
            from reportlab.lib import colors
        except ImportError:
            return Response(
                {'erreur': 'reportlab non disponible. Installez-le avec : pip install reportlab'},
                status=503
            )

        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=A4)
        styles = getSampleStyleSheet()
        elements = []

        # Titre
        elements.append(Paragraph("Rapport Global - Association Coranique", styles['Title']))
        elements.append(Paragraph(f"Généré le {date.today()}", styles['Normal']))
        elements.append(Spacer(1, 0.5 * cm))

        # Données globales
        data = [
            ['Indicateur', 'Valeur'],
            ['Nombre de centres', str(Centre.objects.count())],
            ['Nombre d\'enseignants', str(CustomUser.objects.filter(role='enseignant').count())],
            ['Nombre d\'apprenants', str(CustomUser.objects.filter(role='apprenant').count())],
            ['Nombre de créneaux', str(Creneau.objects.count())],
            ['Nombre de séances de lecture', str(SeanceLecture.objects.count())],
        ]

        table = Table(data, colWidths=[10 * cm, 6 * cm])
        table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#2E4057')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.grey),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F0F4F8')]),
        ]))
        elements.append(table)

        doc.build(elements)
        buffer.seek(0)

        response = HttpResponse(buffer, content_type='application/pdf')
        response['Content-Disposition'] = f'attachment; filename="rapport_global_{date.today()}.pdf"'
        return response


class ExportExcelView(APIView):
    """
    GET /api/statistiques/export-excel/
    Export Excel du rapport global avec openpyxl.
    """
    permission_classes = [EstAdmin]

    def get(self, request):
        try:
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment
        except ImportError:
            return Response(
                {'erreur': 'openpyxl non disponible. Installez-le avec : pip install openpyxl'},
                status=503
            )

        wb = openpyxl.Workbook()

        # === Feuille 1 : Statistiques globales ===
        ws = wb.active
        ws.title = "Global"

        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill("solid", fgColor="2E4057")

        headers = ["Indicateur", "Valeur"]
        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center')

        rows = [
            ("Nombre de centres", Centre.objects.count()),
            ("Nombre d'enseignants", CustomUser.objects.filter(role='enseignant').count()),
            ("Nombre d'apprenants", CustomUser.objects.filter(role='apprenant').count()),
            ("Nombre de créneaux", Creneau.objects.count()),
            ("Nombre de séances de lecture", SeanceLecture.objects.count()),
        ]
        for i, (label, val) in enumerate(rows, 2):
            ws.cell(row=i, column=1, value=label)
            ws.cell(row=i, column=2, value=val)

        ws.column_dimensions['A'].width = 35
        ws.column_dimensions['B'].width = 15

        # === Feuille 2 : Liste des centres avec taux de présence ===
        ws2 = wb.create_sheet("Centres")
        ws2.append(["ID", "Nom", "Ville", "Enseignants", "Apprenants", "Taux Présence Moyen (%)"])
        for centre in Centre.objects.all():
            total_p = Presence.objects.filter(creneau__centre=centre).count()
            presents_p = Presence.objects.filter(creneau__centre=centre, statut='present').count()
            taux_moy = round(presents_p / total_p * 100, 1) if total_p > 0 else 0
            ws2.append([
                centre.pk,
                centre.nom,
                centre.ville,
                CustomUser.objects.filter(centre=centre, role='enseignant').count(),
                CustomUser.objects.filter(centre=centre, role='apprenant').count(),
                taux_moy,
            ])

        # === Feuille 3 : 100 dernières séances de lecture ===
        ws3 = wb.create_sheet("Seances Lecture")
        ws3.append(["ID", "Apprenant", "Enseignant", "Sourate", "Type", "Évaluation", "Date"])
        seances_100 = SeanceLecture.objects.select_related(
            'apprenant', 'enseignant'
        ).order_by('-date_seance')[:100]
        for s in seances_100:
            ws3.append([
                s.pk,
                str(s.apprenant),
                str(s.enseignant),
                s.sourate,
                s.type_lecture,
                s.evaluation,
                str(s.date_seance),
            ])

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)

        response = HttpResponse(
            buffer,
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = f'attachment; filename="rapport_global_{date.today()}.xlsx"'
        return response
